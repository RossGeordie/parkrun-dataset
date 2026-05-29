#!/usr/bin/env python3
"""Complete remaining ETL: locations + volunteers."""
import openpyxl, re, time, psycopg2
from pathlib import Path

T0 = time.time()
def log(m): print(f"[{time.time()-T0:.1f}s] {m}", flush=True)

conn = psycopg2.connect(host="10.21.63.200", dbname="parkrun", user="postgres", password="TempPass123!")
cur = conn.cursor()

# Check existing tables
cur.execute("SELECT table_name FROM information_schema.tables WHERE table_schema='public' ORDER BY table_name")
tables = [r[0] for r in cur.fetchall()]
log(f"Tables: {tables}")

# Step 1: Create locations table from tracked_events if needed
if 'locations' not in tables:
    log("Creating locations table from tracked_events...")
    cur.execute("""CREATE TABLE locations AS 
        SELECT (ROW_NUMBER() OVER ()) AS id,
               event_name, folder, URL, latitude, longitude,
               local_authority, junior, latest_event, guid
        FROM tracked_events""")
    cur.execute("ALTER TABLE locations ADD PRIMARY KEY (id)")
    conn.commit()
    
    # Verify
    cur.execute("SELECT count(*) FROM locations")
    log(f"Locations created: {cur.fetchone()[0]} rows")
    
    # Update events with location_id
    cur.execute("""UPDATE events e SET location_id = l.id
        FROM locations l
        WHERE lower(e.location_folder) = lower(l.folder)""")
    log(f"Events updated with location_id: {cur.rowcount}")
else:
    log("Locations table already exists")
    # Verify events have location_ids
    cur.execute("SELECT count(*) FROM events WHERE location_id IS NULL")
    null_locs = cur.fetchone()[0]
    log(f"Events without location_id: {null_locs}")

# Build folder → location_id map
cur.execute("SELECT id, folder FROM locations")
folder_map = {f[1].lower(): f[0] for f in cur.fetchall()}
log(f"Folder map: {len(folder_map)} entries")

# Step 2: Load volunteers
vol_files = sorted(Path("/home/oc/Documents/Parkrun/Raw2").glob("*-Volunteers-*.xlsx"))
log(f"Volunteer files: {len(vol_files)}")

# Get event_id map
cur.execute("SELECT id, event_num, location_id FROM events WHERE location_id IS NOT NULL")
evt_map = {}
for r in cur.fetchall():
    evt_map[(r[1], r[2])] = r[0]
log(f"Event ID map: {len(evt_map)} entries")

total_vols = 0
skipped = 0

for i, vf in enumerate(vol_files):
    m = re.search(r'Volunteers-(\d+)', vf.stem)
    if not m:
        skipped += 1
        continue
    
    event_num = int(m.group(1))
    loc_folder = vf.stem.split('-Volunteers-')[0]
    
    # Find location
    location_id = 0
    for fl, lid in folder_map.items():
        if loc_folder.lower() == fl or loc_folder.lower().replace(' ', '') == fl.replace(' ', ''):
            location_id = lid
            break
    
    # Find event_id
    eid = evt_map.get((event_num, location_id))
    if not eid:
        skipped += 1
        continue
    
    # Read volunteers
    wb = openpyxl.load_workbook(vf, data_only=True)
    ws = wb.active
    vb_rows = list(ws.iter_rows(values_only=True))
    
    file_vols = []
    for row in vb_rows:
        if not row or not str(row[0]).strip():
            continue
        
        name = str(row[0]).strip()
        pid = None
        if row[1] and str(row[1]).strip():
            m2 = re.search(r'/parkrunner/(\d+)', str(row[1]))
            pid = int(m2.group(1)) if m2 else None
        
        role = None
        if row[3] and str(row[3]).strip():
            role = str(row[3]).strip().split('\n')[0].strip()
        
        club = None
        if row[4] and str(row[4]).strip():
            club = str(row[4]).strip()
        
        file_vols.append((eid, event_num, location_id, pid, name, role, club))
    
    if file_vols:
        cur.executemany("""INSERT INTO volunteers (event_id, event_num, location_id, event_date,
                          parkrun_id, person_name, volunteer_role, club)
                          VALUES (%s,%s,%s,%s,%s,%s,%s,%s)""", file_vols)
        conn.commit()
        total_vols += len(file_vols)

conn.commit()

# Final verification
log("\n=== FINAL COUNTS ===")
for tbl in ['tracked_events', 'events', 'people', 'participation', 'volunteers', 'locations']:
    cur.execute(f"SELECT count(*) FROM {tbl}")
    log(f"  {tbl:20} = {cur.fetchone()[0]:>10,}")

conn.close()
log(f"\nDone in {time.time()-T0:.1f}s")
