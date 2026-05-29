#!/usr/bin/env python3
"""Complete Parkrun ETL - clean run from scratch."""
import psycopg2
import openpyxl
import re
import time
import csv
import glob
from pathlib import Path
from datetime import date

START = time.time()
def log(msg):
    print(f"[{time.time()-START:.1f}s] {msg}", flush=True)

log("=== PARKRUN ETL START ===")

# Create fresh connection (never reuse a broken connection)
conn = psycopg2.connect(host="10.21.63.200", dbname="parkrun", user="postgres", password="TempPass123!")
conn.autocommit = False
cur = conn.cursor()

# === PHASE 0: Schema ===
log("Phase 0: Drop and recreate tables")
for tbl in ['tracked_events', 'events', 'participation', 'volunteers', 'people', 'locations']:
    cur.execute(f"DROP TABLE IF EXISTS {tbl} CASCADE")
conn.commit()

cur.execute("""CREATE TABLE tracked_events (
    id SERIAL PRIMARY KEY, event_name TEXT,
    folder TEXT, latitude REAL, longitude REAL,
    local_authority TEXT, junior TEXT, latest_event INTEGER, guid TEXT
)""")
cur.execute("""CREATE TABLE locations (
    id SERIAL PRIMARY KEY, event_name TEXT,
    folder TEXT, latitude REAL, longitude REAL,
    local_authority TEXT, junior TEXT, latest_event INTEGER, guid TEXT
)""")
cur.execute("""CREATE TABLE events (
    id SERIAL PRIMARY KEY,
    event_num INTEGER NOT NULL,
    location_id INTEGER REFERENCES locations(id),
    event_date DATE NOT NULL,
    UNIQUE(event_num, location_id)
)""")
cur.execute("""CREATE TABLE people (
    id SERIAL PRIMARY KEY,
    person_id INTEGER UNIQUE,
    person_name TEXT,
    gender TEXT,
    total_parkruns INTEGER DEFAULT 0,
    has_100_club BOOLEAN DEFAULT FALSE
)""")
cur.execute("""CREATE TABLE participation (
    id SERIAL PRIMARY KEY,
    event_id INTEGER REFERENCES events(id),
    event_num INTEGER NOT NULL,
    location_id INTEGER NOT NULL,
    event_date DATE NOT NULL,
    finish_position INTEGER,
    parkrun_id INTEGER,
    person_name TEXT,
    gender TEXT,
    finish_time TEXT,
    club TEXT,
    age_grade REAL
)""")
cur.execute("""CREATE TABLE volunteers (
    id SERIAL PRIMARY KEY,
    event_id INTEGER REFERENCES events(id),
    event_num INTEGER NOT NULL,
    location_id INTEGER NOT NULL,
    event_date DATE,
    parkrun_id INTEGER,
    person_name TEXT,
    volunteer_role TEXT,
    club TEXT
)""")
conn.commit()
log("Schema OK")

# === PHASE 1: Load tracked_events from CSV ===
log("Phase 1: TrackedEvents from CSV")
csv_path = Path("/home/oc/Documents/Parkrun/TrackedEvents.csv")
with open(csv_path) as f:
    csv_rows = list(csv.DictReader(f))
log(f"Loaded {len(csv_rows)} rows from CSV")

for i, row in enumerate(csv_rows):
    lat = None
    lon = None
    lat_s = str(row.get("Latitude", "")).strip()
    lon_s = str(row.get("Longitude", "")).strip()
    if lat_s and lat_s not in ("None", ""):
        lat = float(lat_s)
    if lon_s and lon_s not in ("None", ""):
        lon = float(lon_s)
    
    le_str = str(row.get("LatestEventNumber", "")).replace(",", "").strip()
    latest = int(le_str) if le_str and le_str.isdigit() else None
    
    cur.execute("""INSERT INTO tracked_events (event_name, folder, latitude, longitude, local_authority, junior, latest_event, guid)
        VALUES (%s,%s,%s,%s,%s,%s,%s,%s)""",
        (row["Title"], row["folder"], lat, lon, row.get("Local Authoritie"), row.get("Junior","5k"), latest, row["GUID"]))
conn.commit()
log(f"Inserted {len(csv_rows)} tracked_events")

# Build locations table from tracked_events
cur.execute("DELETE FROM locations")
cur.execute("""INSERT INTO locations (id, event_name, folder, latitude, longitude, local_authority, junior, latest_event, guid)
    SELECT id, event_name, folder, latitude, longitude, local_authority, junior, latest_event, guid
    FROM tracked_events ORDER BY id""")
conn.commit()
cur.execute("SELECT setval('locations_id_seq', (SELECT MAX(id) FROM locations))")
conn.commit()
log(f"Locations populated: {cur.fetchone()[0]} rows")

# === PHASE 2: Build mappings ===
log("Phase 2: Build mappings")
cur.execute("SELECT id, folder FROM locations ORDER BY id")
folder_map = {}
for loc_id, folder in cur.fetchall():
    folder_map[folder.lower()] = loc_id
    folder_map[folder.replace(' ', '').lower()] = loc_id
    folder_map[folder.replace('-', ' ').lower()] = loc_id
log(f"Folder map: {len(folder_map)} entries")

# === PHASE 3: Load event dates ===
log("Phase 3: Event dates")
event_dates = {}
raw2 = Path("/home/oc/Documents/Parkrun/Raw2")

history_files = sorted(glob.glob(str(raw2 / "*-event_history.xlsx")))
log(f"History files: {len(history_files)}")

for hf in history_files:
    wb = openpyxl.load_workbook(hf, data_only=True)
    ws = wb.active
    loc_folder = Path(hf).stem.replace("-event_history", "")
    
    # Find location
    loc_id = folder_map.get(loc_folder.lower())
    if not loc_id:
        for fl, lid in folder_map.items():
            if fl.replace(' ', '') == loc_folder.lower().replace(' ', ''):
                loc_id = lid
                break
    if not loc_id:
        continue
    
    for row in ws.iter_rows(values_only=True):
        if not row or row[0] is None:
            continue
        e_num = row[0]
        if not isinstance(e_num, (int, float)):
            continue
        e_num = int(e_num)
        e_date_s = str(row[1]).strip().split('\n')[0].strip() if row[1] else None
        if not e_date_s:
            continue
        m = re.match(r'(\d{2})/(\d{2})/(\d{4})', e_date_s)
        if not m:
            continue
        e_date = date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
        event_dates[(e_num, loc_id)] = e_date

log(f"Dates loaded: {len(event_dates)} entries")

# === PHASE 4: Load events & participation ===
log("Phase 4: Events & participation")
event_files = sorted(glob.glob(str(raw2 / "*-Event-*.xlsx")))
log(f"Event files: {len(event_files)}")

all_events = []
people_dict = {}
part_count = 0
file_count = 0

for ef in event_files:
    # Extract event number: filename = "Town Moor-Event-269.xlsx"
    m = re.search(r'Event-(\d+)\.xlsx$', Path(ef).name)
    if not m:
        continue
    event_num = int(m.group(1))
    
    # Extract location folder
    loc_folder = Path(ef).stem.split('-Event-')[0]
    loc_id = folder_map.get(loc_folder.lower())
    if not loc_id:
        for fl, lid in folder_map.items():
            if fl.replace(' ', '') == loc_folder.lower().replace(' ', ''):
                loc_id = lid
                break
    if not loc_id:
        continue
    
    # Get date
    event_date = event_dates.get((event_num, loc_id))
    if not event_date:
        continue
    
    # Insert event (get event_id)
    cur.execute("""INSERT INTO events (event_num, location_id, event_date)
        VALUES (%s,%s,%s)
        ON CONFLICT (event_num, location_id) DO UPDATE SET event_date=EXCLUDED.event_date
        RETURNING id""", (event_num, loc_id, event_date))
    event_id = cur.fetchone()[0]
    all_events.append((event_id, event_num, loc_id, event_date))
    
    # Read participation from xlsx
    wb = openpyxl.load_workbook(ef, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    
    if len(rows) < 2:
        continue
    
    # Skip header row (index 0), process from row 1
    batch = []
    for row in rows[1:]:
        if not row or len(row) < 11:
            continue
        
        # Column mapping (from prior ETL work):
        # Col 0: chip number
        # Col 1: position
        # Col 2: name
        # Col 3: parkrun URL
        # Col 4: chip number (duplicate)
        # Col 5: gender
        # Col 6: age
        # Col 7: position (gender split)
        # Col 8: age grade
        # Col 9: club
        # Col 10: finish time
        
        finish_position = row[1]
        if not isinstance(finish_position, (int, float)) and not str(finish_position).isdigit():
            continue
        
        parkrun_id = None
        if row[3]:
            m2 = re.search(r'/parkrunner/(\d+)', str(row[3]))
            if m2:
                parkrun_id = int(m2.group(1))
        
        gender = str(row[5]).strip().upper() if row[5] is not None and str(row[5]).strip() else None
        if gender in ("MALE", "M"): gender = "M"
        elif gender in ("FEMALE", "F"): gender = "F"
        
        club = str(row[9]).strip() if row[9] is not None and str(row[9]).strip() else None
        finish_time = str(row[10]).strip() if row[10] is not None else None
        
        batch.append((
            event_id, event_num, loc_id, event_date,
            finish_position, parkrun_id,
            str(row[2]).strip() if row[2] else None,
            gender, finish_time, club, None
        ))
        
        if parkrun_id and parkrun_id not in people_dict:
            people_dict[parkrun_id] = {
                'name': str(row[2]).strip() if row[2] else None,
                'gender': gender
            }
    
    if batch:
        cur.executemany("""INSERT INTO participation (
            event_id, event_num, location_id, event_date,
            finish_position, parkrun_id, person_name, gender, finish_time, club, age_grade
        ) VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)""", batch)
        part_count += len(batch)
    
    file_count += 1
    
    if file_count % 1000 == 0:
        conn.commit()
        log(f"  {file_count}/{len(event_files)} files, {part_count:,} rows")

conn.commit()
log(f"Events: {len(all_events)}, Participation: {part_count:,}")

# === PHASE 5: Load volunteers ===
log("Phase 5: Volunteers")
vol_files = sorted(glob.glob(str(raw2 / "*-Volunteers-*.xlsx")))
log(f"Volunteer files: {len(vol_files)}")

vol_count = 0
vol_proc = 0

for vf in vol_files:
    m = re.search(r'Volunteers-(\d+)\.xlsx$', Path(vf).name)
    if not m:
        continue
    event_num = int(m.group(1))
    
    loc_folder = Path(vf).stem.split('-Volunteers-')[0]
    loc_id = folder_map.get(loc_folder.lower())
    if not loc_id:
        for fl, lid in folder_map.items():
            if fl.replace(' ', '') == loc_folder.lower().replace(' ', ''):
                loc_id = lid
                break
    if not loc_id:
        continue
    
    event_id = None
    for eid, enum, eid_loc, edate in all_events:
        if enum == event_num and eid_loc == loc_id:
            event_id = eid
            break
    
    if not event_id:
        continue
    
    wb = openpyxl.load_workbook(vf, data_only=True)
    ws = wb.active
    vb_rows = list(ws.iter_rows(values_only=True))
    
    vol_batch = []
    for row in vb_rows[1:]:
        if not row or not str(row[0]).strip():
            continue
        
        name = str(row[0]).strip()
        parkrun_id = None
        if row[1] and str(row[1]).strip():
            m2 = re.search(r'/parkrunner/(\d+)', str(row[1]))
            if m2:
                parkrun_id = int(m2.group(1))
        
        role = None
        if row[3] and str(row[3]).strip():
            role = str(row[3]).strip().split('\n')[0].strip()
        
        club = str(row[4]).strip() if row[4] and str(row[4]).strip() else None
        
        vol_batch.append((event_id, event_num, loc_id, event_date, parkrun_id, name, role, club))
    
    if vol_batch:
        cur.executemany("""INSERT INTO volunteers (
            event_id, event_num, location_id, event_date,
            parkrun_id, person_name, volunteer_role, club
        ) VALUES (%s,%s,%s,%s,%s,%s,%s,%s)""", vol_batch)
        vol_count += len(vol_batch)
    
    vol_proc += 1
    if vol_proc % 1000 == 0:
        conn.commit()
        log(f"  {vol_proc}/{len(vol_files)} files, {vol_count:,} volunteers")

conn.commit()
cur.execute("SELECT COUNT(*) FROM volunteers")
log(f"Volunteers: {vol_count} rows (DB: {cur.fetchone()[0]} rows)")

# === PHASE 6: Extract people ===
log("Phase 6: People")
cur.execute("DELETE FROM people")
conn.commit()

# Insert distinct people from participation
cur.execute("""INSERT INTO people (person_id, person_name, gender)
    SELECT DISTINCT ON (parkrun_id) parkrun_id, person_name, gender
    FROM participation
    WHERE parkrun_id IS NOT NULL
    ORDER BY parkrun_id, event_date DESC""")
conn.commit()

# Update totals
cur.execute("""UPDATE people SET total_parkruns = cnt.cnt, has_100_club = cnt.cnt >= 100
FROM (
    SELECT parkrun_id, COUNT(*) as cnt
    FROM participation
    WHERE finish_position IS NOT NULL
    GROUP BY parkrun_id
) cnt WHERE people.person_id = cnt.parkrun_id""")
conn.commit()

cur.execute("SELECT count(*), max(total_parkruns), count(*) FILTER (WHERE has_100_club) FROM people")
r = cur.fetchone()
log(f"People: {r[0]}, max parkruns: {r[1]}, 100+ club: {r[2]}")

# === FINAL STATS ===
log("\n" + "="*60)
log("FINAL RESULTS")
log("="*60)

cur.execute("SELECT count(*) FROM tracking_events")
for tbl in ['tracked_events', 'events', 'people', 'participation', 'volunteers', 'locations']:
    cur.execute(f"SELECT count(*) FROM {tbl}")
    log(f"  {tbl:20} = {cur.fetchone()[0]:>10,}")

log(f"\nTotal time: {time.time()-START:.1f}s")
log("ETL COMPLETE!")

conn.close()
